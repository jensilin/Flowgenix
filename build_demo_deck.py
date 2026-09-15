"""Build the Flow Studio demo deck on top of the Nokia PPT toolkit template.

The template supplies the master, theme fonts and layouts; this script strips the
toolkit's own 83 guidance slides and writes the demo narrative into its layouts.

    python build_demo_deck.py

Inputs:  Presentation2.pptx (Nokia toolkit), deck-assets/*.png (app screenshots)
Output:  NiFi-Flow-Studio-Demo.pptx
"""

from __future__ import annotations

import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

TEMPLATE = Path.home() / "Downloads" / "Presentation2.pptx"
ASSETS = Path("deck-assets")
OUTPUT = Path.home() / "Downloads" / "NiFi-Flow-Studio-Demo.pptx"
WORK = Path.home() / "Downloads" / "_demo-deck-work.pptx"

FOOTER = "NiFi Flow Studio | Demo"

# Layout indices in the Nokia toolkit master.
L_TITLE = 41          # "3 K Blue"  - branded title slide
L_DIVIDER = 59        # "Divider Blue"
L_HEAD = 1            # "1.2 Title" - headline + subtitle, empty canvas below
L_TEXT = 2            # "1.3 Text"
L_COL2 = 4            # "1.5 Bulletpoint text 2 col"
L_COL3 = 5            # "1.6 Bulletpoint text 3 col"
L_NUMBERED = 7        # "1.8 Numbered text"
L_END = 71            # "End slide Blue"

# Screenshots are captured at devicePixelRatio 2.44 against a 1024pt viewport,
# so crops are expressed in viewport points and scaled up.
SCALE = 2497 / 1024

CROPS = {
    "01-landing.png": (295, 18, 892, 600),
    "02-detect.png": (315, 170, 880, 600),
    "03-report.png": (318, 33, 880, 443),
    "04-artifacts.png": (318, 262, 880, 400),
}


def crop_shot(name: str) -> Path:
    """Crop a raw screenshot down to its content card and cache the result."""
    src = ASSETS / name
    out = ASSETS / f"crop-{name}"
    left, top, right, bottom = CROPS[name]
    box = tuple(round(v * SCALE) for v in (left, top, right, bottom))
    with Image.open(src) as im:
        im.crop(box).save(out)
    return out


def trim_template_slides(work: Path, final: Path, drop_count: int) -> None:
    """Delete the toolkit's own slides, using PowerPoint to do the deletion.

    Removing them with python-pptx looks like it works - the package unzips and
    parses cleanly - but PowerPoint still refuses the result as corrupt. The
    template indexes its slides in several places besides the slide id list:
    named sections held in a presentation extension, viewProps recording the
    slides last open in the editor, and a slide count and title vector in
    docProps/app.xml. A stale entry in any one of them is enough. Rather than
    chase every index, the demo slides are appended to the intact template and
    PowerPoint deletes the leading template slides, keeping all of its own
    bookkeeping consistent.
    """
    script = f"""
$ErrorActionPreference = 'Stop'
$pp = New-Object -ComObject PowerPoint.Application
$deck = $pp.Presentations.Open('{work}', $false, $false, $false)
for ($i = {drop_count}; $i -ge 1; $i--) {{ $deck.Slides.Item($i).Delete() }}
$deck.SaveAs('{final}')
$deck.Close()
$pp.Quit()
"""
    subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
        check=True,
        capture_output=True,
        text=True,
    )


def add(prs: Presentation, layout: int):
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 3:
            ph.text_frame.text = FOOTER
    return slide


def ph(slide, idx):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    raise KeyError(f"no placeholder {idx} on this layout")


def write(shape, lines, size=None, lead_bold=False):
    """Fill a placeholder. `lines` may hold (text, indent_level) tuples."""
    tf = shape.text_frame
    tf.word_wrap = True
    normalised = [(l, 0) if isinstance(l, str) else l for l in lines]
    for i, (text, level) in enumerate(normalised):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.level = level
        run = para.add_run()
        run.text = text
        if size is not None:
            run.font.size = Pt(size)
        if lead_bold and i == 0:
            run.font.bold = True


def drop_empty(slide) -> None:
    """Delete placeholders left unfilled so they don't show prompt text."""
    for shape in list(slide.placeholders):
        if not shape.text_frame.text.strip():
            shape._element.getparent().remove(shape._element)


def add_picture(slide, path: Path, top: float, height: float, left: float | None = None):
    with Image.open(path) as im:
        aspect = im.width / im.height
    width = height * aspect
    if left is None:
        left = (10.0 - width) / 2
    pic = slide.shapes.add_picture(
        str(path), Inches(left), Inches(top), Inches(width), Inches(height)
    )
    pic.line.width = Pt(0.75)
    pic.line.color.rgb = RGBColor(0xD3, 0xD8, 0xE0)
    return pic


def caption(slide, text: str, top: float, left: float, width: float):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(0.3)
    )
    tf = box.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(9)
    run.font.italic = True
    return box


def build() -> None:
    prs = Presentation(str(TEMPLATE))
    template_slides = len(prs.slides._sldIdLst)
    shots = {name: crop_shot(name) for name in CROPS}

    # 1 ── Title -----------------------------------------------------------
    s = add(prs, L_TITLE)
    write(ph(s, 0), ["NiFi Flow Studio"])
    write(
        ph(s, 12),
        ["Prompt-driven flow creation and NiFi version migration", "Ravi Raghavendra G S"],
        size=12,
    )

    # 2 ── Divider ---------------------------------------------------------
    s = add(prs, L_DIVIDER)
    write(ph(s, 0), ["Why we built it"])

    # 3 ── Problem / result ------------------------------------------------
    s = add(prs, L_COL2)
    write(ph(s, 0), ["Building and upgrading NiFi flows is slow, manual work"])
    write(ph(s, 10), ["Two problems, one tool"], size=13)
    write(
        ph(s, 15),
        [
            "Today",
            ("A new flow means dragging processors onto the canvas one at a time", 1),
            ("Property names and NAR bundle versions differ per NiFi release, so a flow that works on one instance is invalid on the next", 1),
            ("Upgrading 1.x to 2.x is a manual audit: which processors were removed, which properties were renamed, what breaks silently", 1),
            ("Nothing tells you a flow deployed cleanly but never moved a single FlowFile", 1),
        ],
        size=11,
        lead_bold=True,
    )
    write(
        ph(s, 16),
        [
            "With Flow Studio",
            ("Describe the flow in plain English; it is built, wired and deployed against your live instance", 1),
            ("Component types and bundle versions are read from the running NiFi, never hardcoded", 1),
            ("A 1.x export is analysed component by component against the target release and a migrated flow is generated", 1),
            ("Deployed flows are tested over the REST API and compared against a known-good output file", 1),
        ],
        size=11,
        lead_bold=True,
    )

    # 4 ── Product shot ----------------------------------------------------
    s = add(prs, L_HEAD)
    write(ph(s, 0), ["One console for the whole flow lifecycle"])
    write(
        ph(s, 10),
        ["Connect, build, review, deploy, test and migrate from a single page"],
        size=13,
    )
    add_picture(s, shots["01-landing.png"], top=1.32, height=3.72)
    drop_empty(s)

    # 5 ── Divider ---------------------------------------------------------
    s = add(prs, L_DIVIDER)
    write(ph(s, 0), ["Capability one\nPrompt to running flow"])

    # 6 ── How it works ---------------------------------------------------
    s = add(prs, L_COL3)
    write(ph(s, 0), ["How a sentence becomes a running flow"])
    write(ph(s, 10), ["Every stage is grounded in the NiFi instance you are pointed at"], size=13)
    write(
        ph(s, 15),
        [
            "1. Understand",
            ("The prompt is read against the catalog of components that release actually ships", 1),
            ("On the 1.25.0 instance that is 359 processors and 123 controller services, not a hardcoded shortlist", 1),
        ],
        size=11,
        lead_bold=True,
    )
    write(
        ph(s, 16),
        [
            "2. Build",
            ("A flow spec is produced: process groups, ports, funnels, labels, controller services and connections", 1),
            ("Positions are computed as a left-to-right layered graph, so the canvas is readable", 1),
            ("Unused relationships are auto-terminated before the flow is started", 1),
        ],
        size=11,
        lead_bold=True,
    )
    write(
        ph(s, 17),
        [
            "3. Verify",
            ("Review mode shows the deploy plan and diff without touching NiFi", 1),
            ("After deploy, checks run over the REST API to confirm data actually moved", 1),
            ("A golden file compares the flow's real output against a file you know is correct", 1),
        ],
        size=11,
        lead_bold=True,
    )

    # 7 ── Grounded in the instance ---------------------------------------
    s = add(prs, L_TEXT)
    write(ph(s, 0), ["Nothing about the target version is hardcoded"])
    write(ph(s, 10), ["Point the tool at a different NiFi and the same code adapts"], size=13)
    write(
        ph(s, 15),
        [
            "The version, the processor and controller-service catalog, the NAR bundle versions and the set of supported tuning fields are all read from the live instance at build time.",
            "",
            ("Property names are normalised against each component's own descriptors, so a spec can be written with display names or internal names and still deploy", 1),
            ("NiFi renames properties between releases - Jolt Transformation DSL in 1.x became Jolt Transform in 2.x - and known renames are mapped so one spec deploys to either line", 1),
            ("A property the component does not declare is dropped and reported as a warning rather than silently ignored", 1),
            ("Verified end to end against NiFi 1.25.0 and 2.11.0", 1),
        ],
        size=11,
    )

    # 8 ── Divider ---------------------------------------------------------
    s = add(prs, L_DIVIDER)
    write(ph(s, 0), ["Capability two\nNiFi version migration"])

    # 9 ── Workflow --------------------------------------------------------
    s = add(prs, L_NUMBERED)
    write(ph(s, 0), ["The migration workflow"])
    write(ph(s, 10), ["Six steps, and the original upload is never modified"], size=13)
    write(
        ph(s, 14),
        [
            "Upload a NiFi 1.x template (.xml) or a flow definition (.json) from either line",
            "The source version is detected from the file, or set by hand",
            "Choose the target version from the supported registry - 1.16.3 through 2.11.0",
            "Analyse: every component is classified against the target release",
            "Review the compatibility report, including flow-level concerns",
            "Generate and download the migrated flow plus the report in Markdown and JSON",
        ],
        size=13,
    )

    # 10 ── Detection screenshot ------------------------------------------
    s = add(prs, L_HEAD)
    write(ph(s, 0), ["The source version is read from the file, not guessed"])
    write(
        ph(s, 10),
        ["A .json upload is never assumed to be 2.x - the evidence is shown alongside the verdict"],
        size=13,
    )
    add_picture(s, shots["02-detect.png"], top=1.30, height=3.62)
    caption(
        s,
        "ERPUpdated.xml - detected as NiFi 1.25.0 from the NAR bundle version stamped on 9 of 9 components; 18 components parsed",
        top=5.00,
        left=0.46,
        width=9.09,
    )
    drop_empty(s)

    # 11 ── Report screenshot ---------------------------------------------
    s = add(prs, L_HEAD)
    write(ph(s, 0), ["The compatibility report"])
    write(
        ph(s, 10),
        ["Per-component verdicts, plus the flow-level concerns no single component can tell you about"],
        size=13,
    )
    add_picture(s, shots["03-report.png"], top=1.30, height=3.62)
    caption(
        s,
        "XML templates were removed in 2.0, the Variable Registry no longer exists, and the 2.x line requires Java 21 - each with the action to take",
        top=5.00,
        left=0.46,
        width=9.09,
    )
    drop_empty(s)

    # 12 ── Precedence -----------------------------------------------------
    s = add(prs, L_COL2)
    write(ph(s, 0), ["How compatibility is decided"])
    write(ph(s, 10), ["Deterministic sources win; the model only fills the gaps"], size=13)
    write(
        ph(s, 15),
        [
            "Precedence, highest first",
            ("The target version's installed component catalog, when a target instance is connected - the authority on what exists", 1),
            ("A curated rule base of known changes, with an explanation for each", 1),
            ("Structural checks, such as EVENT_DRIVEN scheduling that 2.x no longer supports", 1),
            ("The AI model, only where the above cannot settle it", 1),
            ("Otherwise reported honestly as not verified", 1),
        ],
        size=11,
        lead_bold=True,
    )
    write(
        ph(s, 16),
        [
            "What the model may and may not do",
            ("It annotates unresolved components with an explanation and a recommendation", 1),
            ("It cannot change a verdict or a replacement type, and its text is never applied to the generated flow automatically", 1),
            ("Every row in the report names who decided it, so a reviewer can see rule base from model at a glance", 1),
            ("Reporting a component as unverified does not block generation - safe deterministic changes are still applied", 1),
        ],
        size=11,
        lead_bold=True,
    )

    # 13 ── Artifacts screenshot ------------------------------------------
    s = add(prs, L_TEXT)
    write(ph(s, 0), ["Generated flow and reports"])
    write(
        ph(s, 10),
        ["A valid 2.x flow definition, importable on the target canvas, with review markers written in"],
        size=13,
    )
    add_picture(s, shots["04-artifacts.png"], top=1.50, height=1.35)
    # A placeholder inherits its geometry from the layout, so overriding only
    # part of it leaves the rest at zero. Set all four to move it below the
    # screenshot.
    body = ph(s, 15)
    body.left, body.top = Inches(0.46), Inches(3.10)
    body.width, body.height = Inches(9.09), Inches(1.95)
    write(
        body,
        [
            "The output is a NiFi VersionedFlowSnapshot, which NiFi deserialises strictly - one unrecognised field and the whole import is rejected with no detail.",
            ("Process-group hierarchy is reproduced rather than flattened, so connections still resolve", 1),
            ("Connection endpoints are emitted as ConnectableComponent objects, and every endpoint is checked against a real component before the file is written", 1),
            ("Components are imported disabled, so nothing moves data before you have reviewed it", 1),
            ("Provenance is written into the root group comments, never as a top-level key", 1),
        ],
        size=11,
    )

    # 14 ── Worked example -------------------------------------------------
    s = add(prs, L_TEXT)
    write(ph(s, 0), ["Worked example: a real ERP flow, 1.25.0 to 2.11.0"])
    write(ph(s, 10), ["ERPUpdated.xml - 86.4 kB, exported from a production 1.25.0 canvas"], size=13)
    write(
        ph(s, 15),
        [
            ("Parsed into 18 components: 7 processors, 2 controller services, 6 connections and 3 labels", 1),
            ("Source version detected as 1.25.0 with certainty, from the bundle version on every component", 1),
            ("Processor types in play: ConvertRecord, EvaluateJsonPath, GenerateFlowFile, InvokeHTTP, PutFile, UpdateAttribute", 1),
            ("Report returned no blocking issues, with three flow-level concerns to act on and one manual-review marker written into the flow", 1),
            ("All 6 connections resolve to real component identifiers - full referential integrity", 1),
            "",
            "This file is what exposed the bug worth knowing about: it imported as a correct-looking component count with every connection empty, because the parser was reading flat sourceId fields that real NiFi templates do not use. Connection endpoints live in nested source and destination elements. That is now covered by tests.",
        ],
        size=11,
    )

    # 15 ── Confidence -----------------------------------------------------
    s = add(prs, L_COL2)
    write(ph(s, 0), ["Why you can trust the output"])
    write(ph(s, 10), ["The failure mode we designed against is a flow that looks right and is not"], size=13)
    write(
        ph(s, 15),
        [
            "Tested",
            ("140 automated tests across the version registry, the rule base, flow parsing, the migration engine and snapshot generation, plus a smoke suite that runs against a live server", 1),
            ("Structural tests assert the exact top-level keys NiFi accepts, that every component carries its required fields, and that every connection endpoint resolves", 1),
            ("The suite catches an invalid flow before NiFi does, with a precise message instead of \"an unexpected error has occurred\"", 1),
        ],
        size=11,
        lead_bold=True,
    )
    write(
        ph(s, 16),
        [
            "Safe by default",
            ("The uploaded file is never modified - migration works on a copy", 1),
            ("Generated components arrive disabled and are yours to start", 1),
            ("Review mode shows the plan and diff without touching NiFi", 1),
            ("Credentials come from .env or the browser session only, and are never written back", 1),
            ("Anything the tool could not verify is labelled as such rather than presented as a clean result", 1),
        ],
        size=11,
        lead_bold=True,
    )

    # 16 ── Demo script ----------------------------------------------------
    s = add(prs, L_COL3)
    write(ph(s, 0), ["What we will run live"])
    write(ph(s, 10), ["Roughly ten minutes end to end"], size=13)
    write(
        ph(s, 15),
        [
            "Build a flow",
            ("Detect the NiFi version", 1),
            ("Type a prompt for a JSON to CSV flow and pick the columns to keep", 1),
            ("Review the plan, deploy, then watch live status and the flow test result", 1),
        ],
        size=11,
        lead_bold=True,
    )
    write(
        ph(s, 16),
        [
            "Migrate a flow",
            ("Upload ERPUpdated.xml and let it detect 1.25.0", 1),
            ("Target 2.11.0 and analyse", 1),
            ("Walk the report, then generate and import the migrated flow on the 2.11.0 canvas", 1),
        ],
        size=11,
        lead_bold=True,
    )
    write(
        ph(s, 17),
        [
            "Worth asking about",
            ("Pointing it at a different NiFi release", 1),
            ("Adding a rule for a component specific to your estate", 1),
            ("Per-model token usage, tracked across every run", 1),
        ],
        size=11,
        lead_bold=True,
    )

    # 17 ── End ------------------------------------------------------------
    add(prs, L_END)

    count = len(prs.slides._sldIdLst) - template_slides
    prs.save(str(WORK))
    trim_template_slides(WORK, OUTPUT, template_slides)
    WORK.unlink(missing_ok=True)
    print(f"wrote {OUTPUT} ({count} slides)")


if __name__ == "__main__":
    build()
